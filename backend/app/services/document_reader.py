from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from urllib.parse import urlsplit

import pymupdf
from docx import Document
from docx.table import Table


MAX_FILE_SIZE = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class DocumentReadError(ValueError):
    """Lỗi file CV không thể xử lý ở phiên bản hiện tại."""


def _safe_link(value: str) -> bool:
    # Preserve document metadata as text only; never request these URLs.
    try:
        parsed = urlsplit(value)
        return parsed.scheme.lower() in {"https", "http", "mailto", "tel"} and not any(char in value for char in "\n\r")
    except ValueError:
        return False


def _paragraph_text(paragraph) -> str:
    text = paragraph.text.strip()
    links = []
    for link in paragraph.hyperlinks:
        if link.url and _safe_link(link.url) and link.url not in text and link.url not in links:
            links.append(link.url)
    return " ".join([text, *links]).strip()


def validate_cv_file(filename: str, file_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentReadError("Chỉ hỗ trợ file PDF hoặc DOCX")
    if not file_bytes:
        raise DocumentReadError("File CV đang rỗng")
    if len(file_bytes) > MAX_FILE_SIZE:
        raise DocumentReadError("File CV vượt quá giới hạn 5 MB")
    return extension


def read_pdf(file_bytes: bytes) -> str:
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            if document.needs_pass:
                raise DocumentReadError("PDF có mật khẩu; cần bỏ mật khẩu trước khi upload")
            pages = []
            for page in document:
                # Best-effort block order; not a guarantee for arbitrary multi-column layouts.
                blocks = sorted(page.get_text("blocks"), key=lambda block: (block[1], block[0]))
                links = page.get_links()
                lines = []
                for block in blocks:
                    text = block[4].strip()
                    if not text:
                        continue
                    lines.append(text)
                    for link in links:
                        uri = link.get("uri", "")
                        if uri and _safe_link(uri) and pymupdf.Rect(block[:4]).intersects(link["from"]) and uri not in text:
                            label = page.get_textbox(link["from"]).strip().replace("\n", " ")
                            lines.append(f"{label}: {uri}" if label else uri)
                pages.append("\n".join(lines))
    except DocumentReadError:
        raise
    except Exception as error:
        raise DocumentReadError("Không thể đọc file PDF") from error

    text = "\n\n".join(page for page in pages if page).strip()
    if not text:
        raise DocumentReadError("PDF không có lớp văn bản; file này cần OCR")
    return text


def read_docx(file_bytes: bytes) -> str:
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                raise DocumentReadError("DOCX vượt giới hạn nội dung giải nén 25 MB hoặc 1000 thành phần")
        document = Document(BytesIO(file_bytes))
    except DocumentReadError:
        raise
    except Exception as error:
        raise DocumentReadError("Không thể đọc file DOCX") from error

    lines = []
    seen_headers = set()
    for section in document.sections:
        header = "\n".join(filter(None, (_paragraph_text(p) for p in section.header.paragraphs)))
        if header and header not in seen_headers:
            lines.append(header)
            seen_headers.add(header)
    # Preserve interleaved body paragraph/table order, rather than appending all tables last.
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                values = ["\n".join(filter(None, (_paragraph_text(p) for p in cell.paragraphs))) for cell in row.cells]
                values = [value for value in values if value]
                if values:
                    lines.append(" | ".join(values))
        else:
            text = _paragraph_text(block)
            if text:
                lines.append(text)
    seen_footers = set()
    for section in document.sections:
        footer = "\n".join(filter(None, (_paragraph_text(p) for p in section.footer.paragraphs)))
        if footer and footer not in seen_footers:
            lines.append(footer)
            seen_footers.add(footer)
    text = "\n".join(line for line in lines if line).strip()
    if not text:
        raise DocumentReadError("DOCX không có nội dung văn bản")
    return text


def read_cv(filename: str, file_bytes: bytes) -> str:
    extension = validate_cv_file(filename, file_bytes)
    if extension == ".pdf":
        return read_pdf(file_bytes)
    return read_docx(file_bytes)

